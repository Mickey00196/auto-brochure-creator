"""§15 One Pager — single-page executive summary, email-attachment sized.

Headline stats + the §16 comparison table, deliberately no per-unit detail
pages. Built as a single slide so the PDF export (via pdf_export.pptx_to_pdf)
is a literal one-page document.
"""
from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches
from sqlalchemy.orm import Session

from app.models import Proposal
from app.services.brochure.slide_kit import (
    SLIDE_HEIGHT,
    SLIDE_WIDTH,
    add_bullets,
    add_text,
    blank_slide,
    fill_background,
    style_table_cell,
    style_table_header,
)
from app.services.brochure.theme import DEFAULT_THEME, Theme
from app.services.comparison import build_comparison_table
from app.services.qa import QASeverity, run_qa_pass


def build_one_pager(db: Session, proposal: Proposal, *, theme: Theme = DEFAULT_THEME, require_qa_pass: bool = True) -> Presentation:
    qa_report = run_qa_pass(db, proposal)
    if require_qa_pass and not qa_report.is_export_ready:
        blocking = [i.message for i in qa_report.issues if i.severity == QASeverity.BLOCKING]
        raise ValueError(f"Proposal is not export-ready: {blocking}")

    units = proposal.selected_units
    rows = build_comparison_table(units)
    ready_rows = [r for r in rows if not r.is_tbd]

    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    slide = blank_slide(prs)
    fill_background(slide, theme.light_bg_hex)

    add_text(slide, Inches(0.6), Inches(0.35), Inches(10), Inches(0.35), "OFFICE SHORTLIST — EXECUTIVE SUMMARY", size=12, color_hex=theme.accent_hex, bold=True, font_family=theme.font_family)
    add_text(slide, Inches(0.6), Inches(0.7), Inches(11.5), Inches(0.6), proposal.title, size=26, color_hex=theme.text_on_light_hex, bold=True, font_family=theme.font_family)
    add_text(slide, Inches(0.6), Inches(1.3), Inches(11.5), Inches(0.4), f"Prepared for {proposal.client.company_name} by {proposal.prepared_by or 'your advisor'}", size=13, color_hex=theme.muted_text_hex, font_family=theme.font_family)

    stats = [
        f"{len(units)} locations shortlisted",
        f"Lowest all-in rate: €{ready_rows[0].all_in_rate_eur_per_m2_year:,.0f}/m²/yr" if ready_rows else "Pricing pending on all units",
        f"Area range: {min(u.available_area_m2 for u in units):,.0f}–{max(u.available_area_m2 for u in units):,.0f} m²",
    ]
    add_bullets(slide, Inches(0.6), Inches(1.85), Inches(11.5), Inches(0.9), stats, size=13, color_hex=theme.text_on_light_hex, font_family=theme.font_family)

    headers = ["Location", "Floor", "Area (m²)", "All-in €/m²/yr", "Est. Annual Cost"]
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(0.6), Inches(2.9), Inches(11.9), Inches(4.2))
    table = table_shape.table
    for c, header in enumerate(headers):
        style_table_header(table.cell(0, c), theme, header)
    for r, row in enumerate(rows, start=1):
        values = [
            row.building_name,
            row.floor or "—",
            f"{row.available_area_m2:,.0f}",
            f"€{row.all_in_rate_eur_per_m2_year:,.0f}" if row.all_in_rate_eur_per_m2_year is not None else "TBD",
            f"€{row.estimated_annual_cost_eur:,.0f}" if row.estimated_annual_cost_eur is not None else "TBD",
        ]
        for c, val in enumerate(values):
            style_table_cell(table.cell(r, c), theme, val)

    return prs
